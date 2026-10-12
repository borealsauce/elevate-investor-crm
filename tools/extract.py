#!/usr/bin/env python3
"""Extract Elevate_Investor_CRM.xlsx -> data.js for the static site.
Usage: python3 tools/extract.py [--include-contacts] path/to/Elevate_Investor_CRM.xlsx
Formulas are recalculated first with LibreOffice (headless) so cached values exist.
By default PHONE / EMAIL columns are stripped (public site)."""
import sys, os, json, subprocess, tempfile, shutil, datetime
import openpyxl

args = [a for a in sys.argv[1:] if not a.startswith('--')]
INCLUDE_CONTACTS = '--include-contacts' in sys.argv
src = args[0] if args else '/workspace/elevate_raise/Elevate_Investor_CRM.xlsx'
here = os.path.dirname(os.path.abspath(__file__))
out_js = os.path.join(here, '..', 'data.js')

tmp = tempfile.mkdtemp()
shutil.copy(src, os.path.join(tmp, 'in.xlsx'))
subprocess.run(['soffice', '--headless', '--convert-to', 'xlsx:Calc MS Excel 2007 XML',
                '--outdir', os.path.join(tmp, 'out'), os.path.join(tmp, 'in.xlsx')],
               check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
wb = openpyxl.load_workbook(os.path.join(tmp, 'out', 'in.xlsx'), data_only=True)
NAV = [ws.title for ws in wb]
PRIVATE = {'PHONE', 'EMAIL'}

def val(v):
    if isinstance(v, (datetime.datetime, datetime.date)):
        return v.strftime('%Y-%m-%d')
    if isinstance(v, float) and v.is_integer():
        return int(v)
    if isinstance(v, str):
        return v.strip()
    return v

def is_header(cells):
    s = [c for c in cells if c is not None]
    return (len(s) >= 2 and all(isinstance(c, str) and c == c.upper() for c in s)
            and sum(any(ch.isalpha() for ch in c) for c in s) >= 2)

def blocks(ws):
    """Split a sheet into title / table blocks (column positions preserved)."""
    out, table = [], None
    for r in ws.iter_rows(min_row=1, values_only=True):
        cells = [val(c) for c in r[1:]]  # column A is a margin
        nonnull = [c for c in cells if c not in (None, '')]
        if not nonnull:
            table = None; continue
        if [c for c in cells if c is not None] == NAV:
            continue
        if is_header(cells) and not (len(nonnull) == 1):
            cols = [i for i, c in enumerate(cells) if c is not None]
            table = {'type': 'table', 'header': [cells[i] for i in cols], 'cols': cols, 'rows': []}
            out.append(table); continue
        if table is not None:
            table['rows'].append([cells[i] if i < len(cells) else None for i in table['cols']])
        elif len(nonnull) == 1:
            out.append({'type': 'title', 'text': str(nonnull[0])})
        else:
            out.append({'type': 'line', 'cells': nonnull})
    # strip private columns
    if not INCLUDE_CONTACTS:
        for b in out:
            if b['type'] == 'table':
                keep = [i for i, h in enumerate(b['header']) if h not in PRIVATE]
                b['header'] = [b['header'][i] for i in keep]
                b['rows'] = [[r[i] for i in keep] for r in b['rows']]
    return out

def rows_of(ws, n):
    return [[val(c) for c in r[1:n+1]] for r in ws.iter_rows(min_row=1, values_only=True)]

data = {'generated': datetime.date.today().isoformat(), 'source': os.path.basename(src),
        'contactsIncluded': INCLUDE_CONTACTS, 'sheets': {}}

# Dashboard: pull labelled KPI pairs (label row followed by value row, then caption row)
d = rows_of(wb['Dashboard'], 15)
kpis = []
for i in range(len(d) - 2):
    lab, v, cap = d[i], d[i+1], d[i+2]
    labs = [c for c in lab if c is not None]
    if len(labs) == 5 and all(isinstance(c, str) and c.isupper() for c in labs) and lab[0] != '#':
        for j, c in enumerate(lab):
            if c is not None and isinstance(c, str):
                kpis.append({'label': c, 'value': v[j], 'caption': cap[j]})
data['dashboard'] = {'title': d[1][0], 'subtitle': d[2][0], 'kpis': kpis}

for ws in wb:
    if ws.title == 'Dashboard':
        continue
    data['sheets'][ws.title] = blocks(ws)

with open(out_js, 'w') as f:
    f.write('// Generated from ' + os.path.basename(src) + ' — do not edit by hand\n')
    f.write('window.CRM = ' + json.dumps(data, ensure_ascii=False, indent=1) + ';\n')
shutil.rmtree(tmp)
print('wrote', out_js, 'kpis', len(kpis), 'sheets', list(data['sheets']))
